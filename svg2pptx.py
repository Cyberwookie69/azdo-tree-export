"""
svg2pptx.py — zet architecture.svg om naar native PPTX-shapes (geen bitmap).

Gericht op EXACT deze SVG; parseert <rect>/<line>/<path>/<text> en projecteert
die 1:1 (SVG px -> slide pt) op een 1400x900pt slide (19.44" x 12.5").

DESIGN DECISIONS
  1. Slide-maat == SVG viewBox in points (1400x900). Fontsizes (22/16/13/11px)
     mappen 1:1 naar pt -> leesbaar in PowerPoint, geen scale-skew.
     Why: fonts & shapes blijven proportioneel. Gebruiker kan slide verkleinen
     bij presentatie zonder dat tekst te klein wordt in bewerkmodus.
  2. CSS-klassen worden gehardcode gemapt naar fill/stroke/font (ipv een echte
     SVG CSS-parser). Why: de SVG is bekend en vast; parser-complexiteit niet
     nodig. Snelste correcte oplossing.
  3. Paths: alleen "Mx,y Lx,y" wordt geparset (dat is alles wat de SVG gebruikt).
     Why: volledige SVG path parser is overkill voor rechte connectors.
  4. Arrow-markers via directe XML op <a:ln>/<a:tailEnd>; python-pptx exposes
     dit niet high-level.
     Why: enige manier om native arrow-ends te zetten zonder PNG-emulatie.
  5. <defs>/<style> worden overgeslagen; children van <g> plat uitgelopen.
     Why: anders zouden marker-paths in <defs> als connectors verschijnen.
"""

import re
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

SVG_PATH  = r"C:\projects\azdo-tree-export\architecture.svg"
PPTX_PATH = r"C:\projects\azdo-tree-export\architecture.pptx"

CLASS_TEXT = {
    'title':  {'font': 22, 'bold': True,  'color': '111111'},
    'panel':  {'font': 16, 'bold': True,  'color': '111111'},
    'boxlbl': {'font': 13, 'bold': False, 'color': '111111'},
    'lbl':    {'font': 12, 'bold': False, 'color': '333333'},
    'small':  {'font': 11, 'bold': False, 'color': '555555'},
    'mono':   {'font': 11, 'bold': False, 'color': '111111', 'mono': True},
}

CLASS_RECT = {
    'box':    {'fill': 'F6F8FA', 'line': '555555', 'lw': 1.5},
    'ext':    {'fill': 'EEF5FF', 'line': '0969DA', 'lw': 1.5},
    'store':  {'fill': 'EAFBEA', 'line': '1A7F37', 'lw': 1.5},
    'user':   {'fill': 'FFF3CD', 'line': '9A6700', 'lw': 1.5},
    'depbox': {'fill': 'FFFBEA', 'line': '9A6700', 'lw': 1.5},
}

CLASS_EDGE = {
    'edge':     {'color': '333333'},
    'edgeblue': {'color': '0969DA'},
    'edgegrn':  {'color': '1A7F37'},
}

ARROW_MARKERS = {'url(#arrow)', 'url(#arrow-blue)', 'url(#arrow-green)'}


def strip_ns(tag):
    return tag.split('}', 1)[1] if '}' in tag else tag


def iter_shapes(parent):
    """Depth-first, overslaand <defs>/<style>; vlakt <g> uit."""
    for child in parent:
        tag = strip_ns(child.tag)
        if tag in ('defs', 'style'):
            continue
        if tag == 'g':
            yield from iter_shapes(child)
        else:
            yield child


def hex_norm(s, default='FFFFFF'):
    if not s or s.lower() == 'none':
        return None
    s = s.lstrip('#').upper()
    if len(s) == 3:
        s = ''.join(c*2 for c in s)
    return s if len(s) == 6 else default


def add_arrow_tailend(line_shape):
    ln = line_shape.line._get_or_add_ln()
    tailEnd = ln.find(qn('a:tailEnd'))
    if tailEnd is None:
        tailEnd = etree.SubElement(ln, qn('a:tailEnd'))
    tailEnd.set('type', 'triangle')
    tailEnd.set('w', 'med')
    tailEnd.set('len', 'med')


def set_dash(line_shape):
    ln = line_shape.line._get_or_add_ln()
    if ln.find(qn('a:prstDash')) is None:
        prstDash = etree.SubElement(ln, qn('a:prstDash'))
        prstDash.set('val', 'dash')


def add_rect(slide, x, y, w, h, rx, fill_hex, line_hex, line_w):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0 else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill_hex:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    else:
        s.fill.background()
    if line_hex:
        s.line.color.rgb = RGBColor.from_string(line_hex)
        s.line.width = Pt(line_w)
    s.text_frame.text = ''
    if rx > 0:
        try:
            s.adjustments[0] = min(0.5, rx / min(w, h))
        except Exception:
            pass
    return s


def add_connector(slide, x1, y1, x2, y2, color, width, arrow=False, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    if dashed:
        set_dash(c)
    if arrow:
        add_arrow_tailend(c)
    return c


def add_text(slide, x, y, text, cls, anchor, bold_override):
    style = CLASS_TEXT.get(cls, CLASS_TEXT['lbl']).copy()
    if bold_override is not None:
        style['bold'] = bold_override
    size = style['font']

    # SVG text y = baseline; schat top-y
    top_y = y - size * 0.95
    box_w = 420
    if anchor == 'middle':
        left_x = x - box_w / 2
        align = PP_ALIGN.CENTER
    elif anchor == 'end':
        left_x = x - box_w
        align = PP_ALIGN.RIGHT
    else:
        left_x = x
        align = PP_ALIGN.LEFT

    tb = slide.shapes.add_textbox(Pt(left_x), Pt(top_y), Pt(box_w), Pt(size*1.6))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = style['bold']
    r.font.color.rgb = RGBColor.from_string(style['color'])
    r.font.name = 'Consolas' if style.get('mono') else 'Segoe UI'
    return tb


def main():
    prs = Presentation()
    prs.slide_width  = Pt(1400)
    prs.slide_height = Pt(900)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    tree = ET.parse(SVG_PATH)
    root = tree.getroot()

    for el in iter_shapes(root):
        tag = strip_ns(el.tag)
        cls = el.get('class', '')

        if tag == 'rect':
            x = float(el.get('x', 0)); y = float(el.get('y', 0))
            w = float(el.get('width', 0)); h = float(el.get('height', 0))
            rx = float(el.get('rx', 0))
            if cls in CLASS_RECT:
                st = CLASS_RECT[cls]
                add_rect(slide, x, y, w, h, rx, st['fill'], st['line'], st['lw'])
            else:
                fill = hex_norm(el.get('fill'), 'FFFFFF')
                line = hex_norm(el.get('stroke'), 'DDDDDD')
                add_rect(slide, x, y, w, h, rx, fill, line, 1.0)

        elif tag == 'text':
            x = float(el.get('x', 0)); y = float(el.get('y', 0))
            text = ''.join(el.itertext())
            anchor = el.get('text-anchor', 'start')
            if anchor not in ('middle', 'end'):
                anchor = 'start'
            fw = el.get('font-weight')
            bold = True if fw in ('700', 'bold') else None
            add_text(slide, x, y, text, cls, anchor, bold)

        elif tag == 'line':
            x1 = float(el.get('x1', 0)); y1 = float(el.get('y1', 0))
            x2 = float(el.get('x2', 0)); y2 = float(el.get('y2', 0))
            stroke = hex_norm(el.get('stroke'), '333333') or '333333'
            dashed = bool(el.get('stroke-dasharray'))
            add_connector(slide, x1, y1, x2, y2, stroke, 1.0, dashed=dashed)

        elif tag == 'path':
            d = el.get('d', '')
            m = re.match(r'\s*M\s*([-\d.]+)\s*[, ]\s*([-\d.]+)\s*L\s*([-\d.]+)\s*[, ]\s*([-\d.]+)', d)
            if not m:
                continue
            x1, y1, x2, y2 = map(float, m.groups())
            arrow = el.get('marker-end', '') in ARROW_MARKERS
            st = CLASS_EDGE.get(cls, {'color': '333333'})
            add_connector(slide, x1, y1, x2, y2, st['color'], 1.5, arrow=arrow)

    prs.save(PPTX_PATH)
    print(f"OK -> {PPTX_PATH}")


if __name__ == '__main__':
    main()
